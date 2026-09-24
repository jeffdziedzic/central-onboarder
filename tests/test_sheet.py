from pathlib import Path

import openpyxl
import pytest

from central_onboarder.core import sheet as sheet_module


def _device(serial, **overrides):
    d = {
        "serial": serial, "mac": "11:22:33:44:AA:BB", "device_type": "AP",
        "target_group": "Building-1-APs", "target_site": "Building 1", "subscription_key": "KEY123",
    }
    d.update(overrides)
    return d


def test_add_devices_manual_creates_new_sheet_from_template(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    report = sheet_module.add_devices_manual(out, [_device("S1"), _device("S2")])
    assert report.rows_added == 2
    assert report.rows_updated == 0
    assert out.exists()

    rows = sheet_module.read_devices(out)
    assert {r.serial for r in rows} == {"S1", "S2"}


def test_add_devices_manual_merges_by_serial(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1", target_group="Old Group")])
    report = sheet_module.add_devices_manual(out, [_device("S1", target_group="New Group")])
    assert report.rows_added == 0
    assert report.rows_updated == 1
    rows = sheet_module.read_devices(out)
    assert len(rows) == 1
    assert rows[0].target_group == "New Group"


def test_add_devices_manual_skips_blank_serial(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    report = sheet_module.add_devices_manual(out, [_device(""), _device("S1")])
    assert report.skipped_blank_serial_rows == 1
    assert report.rows_added == 1


def test_add_devices_manual_never_touches_tracking_columns_on_merge(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1")])
    sheet_module.mark_added_to_glcp(out, ["S1"])
    sheet_module.add_devices_manual(out, [_device("S1", target_group="New Group")])
    rows = sheet_module.read_devices(out)
    assert rows[0].added_to_glcp == "Y"
    assert rows[0].target_group == "New Group"


def test_mark_functions_only_touch_matching_serials(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1"), _device("S2")])
    sheet_module.mark_preprovisioned(out, ["S1"])
    rows = {r.serial: r for r in sheet_module.read_devices(out)}
    assert rows["S1"].preprovisioned == "Y"
    assert rows["S2"].preprovisioned is None


def test_add_devices_from_csv_maps_rows(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    csv_rows = [{"serial": "S1", "mac": "AA:BB:CC:DD:EE:FF", "device_type": "Switch"}]
    report = sheet_module.add_devices_from_csv(csv_rows, out)
    assert report.rows_added == 1
    rows = sheet_module.read_devices(out)
    assert rows[0].device_type == "Switch"


def test_validate_sheet_schema_ok_for_fresh_sheet(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1")])
    result = sheet_module.validate_sheet_schema(out)
    assert result.ok is True
    assert result.mismatches == []


def test_validate_sheet_schema_detects_renamed_header(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1")])
    wb = openpyxl.load_workbook(out)
    wb[sheet_module.DEVICES_SHEET].cell(row=1, column=1).value = "Serial Number"
    wb.save(out)
    result = sheet_module.validate_sheet_schema(out)
    assert result.ok is False
    assert result.mismatches[0].column == 1


def test_expected_hash_guard_aborts_on_concurrent_edit(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1")])
    stale_hash = sheet_module.hash_existing(out)
    sheet_module.add_devices_manual(out, [_device("S2")])  # changes the file on disk

    report = sheet_module.add_devices_manual(out, [_device("S3")], expected_hash=stale_hash)
    assert report.aborted_reason is not None
    rows = sheet_module.read_devices(out)
    assert {r.serial for r in rows} == {"S1", "S2"}  # S3 never written


def test_load_destinations_populates_reference_sheet_and_dropdown(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1")])
    report = sheet_module.load_destinations(out, groups={"Group A": "id1"}, sites={"Site A": "id2"})
    assert report.groups_loaded == 1
    assert report.sites_loaded == 1
    wb = openpyxl.load_workbook(out)
    assert sheet_module.REFERENCE_SHEET in wb.sheetnames
    ref_ws = wb[sheet_module.REFERENCE_SHEET]
    assert ref_ws.cell(row=sheet_module.REF_DATA_START_ROW, column=sheet_module.REF_COL_GROUP).value == "Group A"


def test_grows_past_template_styled_range(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    devices = [_device(f"S{i}") for i in range(sheet_module.TEMPLATE_STYLED_LAST_ROW + 5)]
    report = sheet_module.add_devices_manual(out, devices)
    assert report.rows_added == len(devices)
    rows = sheet_module.read_devices(out)
    assert len(rows) == len(devices)


# --- Hostname / Hostname Set columns (v0.4.0) ------------------------------


def _write_v1_sheet(path: Path, rows: list[list]) -> None:
    """A sheet in the pre-Hostname 12-column layout, with data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_module.DEVICES_SHEET
    ws.append(list(sheet_module._V1_DEVICES_HEADERS))
    for r in rows:
        ws.append(r)
    ws.auto_filter.ref = "A1:L101"
    wb.save(path)


def test_hostname_round_trips_through_add_and_read(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1", hostname="BLDG1-AP-01")])
    assert sheet_module.read_devices(out)[0].hostname == "BLDG1-AP-01"


def test_add_devices_from_csv_maps_hostname(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_from_csv([{"serial": "S1", "hostname": "AP-01"}], out)
    assert sheet_module.read_devices(out)[0].hostname == "AP-01"


def test_mark_hostname_set(tmp_path: Path):
    out = tmp_path / "device-list.xlsx"
    sheet_module.add_devices_manual(out, [_device("S1"), _device("S2")])
    sheet_module.mark_hostname_set(out, ["S2"])
    rows = {r.serial: r for r in sheet_module.read_devices(out)}
    assert rows["S1"].hostname_set is None
    assert rows["S2"].hostname_set == "Y"
    assert rows["S2"].site_assigned is None


def test_upgrade_moves_existing_data_into_new_layout(tmp_path: Path):
    path = tmp_path / "old.xlsx"
    _write_v1_sheet(path, [
        ["S1", "AA:BB", "AP", "G1", "Site1", "KEY", "Y", "Y", "Y", "Y", "Y", "note one"],
        ["S2", "CC:DD", "Switch", None, None, None, "Y", None, None, None, None, None],
    ])
    assert sheet_module.validate_sheet_schema(path).ok is False
    assert sheet_module.upgrade_sheet_schema(path) is True
    assert sheet_module.upgrade_sheet_schema(path) is False  # idempotent
    assert sheet_module.validate_sheet_schema(path).ok is True

    rows = {r.serial: r for r in sheet_module.read_devices(path)}
    s1 = rows["S1"]
    assert (s1.subscription_key, s1.hostname) == ("KEY", None)
    assert (s1.added_to_glcp, s1.subscription_assigned, s1.service_assigned,
            s1.preprovisioned, s1.site_assigned) == ("Y",) * 5
    assert s1.hostname_set is None
    assert s1.notes == "note one"
    assert rows["S2"].added_to_glcp == "Y"

    ws = openpyxl.load_workbook(path)[sheet_module.DEVICES_SHEET]
    assert ws.auto_filter.ref == "A1:N101"


def test_merging_into_old_sheet_upgrades_first_never_corrupts_tracking(tmp_path: Path):
    """Without the upgrade, a new Hostname would be written into the old
    layout's column 7 (Added to GLCP)."""
    path = tmp_path / "old.xlsx"
    _write_v1_sheet(path, [["S1", "AA:BB", "AP", None, None, None, None, None, None, None, None, None]])
    sheet_module.add_devices_manual(path, [_device("S1", hostname="AP-01")])
    row = sheet_module.read_devices(path)[0]
    assert row.hostname == "AP-01"
    assert row.added_to_glcp is None


def test_upgrade_leaves_unknown_layouts_alone(tmp_path: Path):
    path = tmp_path / "weird.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = sheet_module.DEVICES_SHEET
    wb.active.append(["Serial", "Something Else"])
    wb.save(path)
    assert sheet_module.upgrade_sheet_schema(path) is False
    assert sheet_module.validate_sheet_schema(path).ok is False


def test_bundled_template_is_current_layout():
    assert sheet_module.validate_sheet_schema(sheet_module._template_path()).ok is True
